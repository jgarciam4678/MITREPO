"""Importa el Excel de seguimiento semanal (hojas «CWnn») al tablero Flujo.

Uso:
    # Solo diagnóstico (no escribe nada):
    python -m app.importar_seguimiento_cw seguimiento.xlsx

    # Diagnóstico + carga de las tarjetas abiertas en la base de datos de Flujo:
    python -m app.importar_seguimiento_cw seguimiento.xlsx --db flujo.db

Toma la hoja CW más reciente como estado actual. Solo carga lo que sigue abierto
(lo terminado se queda en el Excel como histórico). Lo que no cumple la
Definición de Listo entra en Backlog para revisarlo en la primera reposición;
lo que ya estaba en proceso se mantiene en curso aunque le falten datos.

Columnas esperadas: Semana, Nombre, Actividad, Proyecto Relacionado, Estado,
Fecha, Prioridad, Comentarios.
"""

from __future__ import annotations

import argparse
import re
import sqlite3
from collections import Counter
from datetime import date, datetime, time, timezone

from openpyxl import load_workbook

from app.flujo import ESQUEMA, faltas_listo

ESTADOS_EXCEL = {
    "por iniciar": "listo",
    "en proceso": "en_curso",
    "en pausa": "en_curso",  # entra bloqueada: «en pausa» no existe en el marco
    "terminado": "hecho",
}
VACIOS = {"", "-", "#n/d", "none"}


def _vacio(valor) -> bool:
    return valor is None or str(valor).strip().lower() in VACIOS


def _semana(texto) -> int | None:
    m = re.search(r"(\d{1,2})", str(texto or ""))
    return int(m.group(1)) if m else None


def leer_hoja_actual(ruta: str) -> tuple[str, int, list[dict]]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    hojas = [
        (int(m.group(1)), nombre)
        for nombre in libro.sheetnames
        if (m := re.fullmatch(r"\s*CW\s*(\d{1,2})\s*", nombre, re.I))
        and libro[nombre].sheet_state == "visible"
    ]
    if not hojas:
        raise SystemExit("No hay hojas visibles con nombre CWnn en el Excel.")
    semana, nombre = max(hojas)
    filas = libro[nombre].iter_rows(values_only=True)
    cabecera = [str(c or "").strip().lower() for c in next(filas)]

    def col(*nombres):
        for n in nombres:
            if n in cabecera:
                return cabecera.index(n)
        return None

    idx = {
        "semana": col("semana"),
        "nombre": col("nombre", "responsable"),
        "actividad": col("actividad", "descripcion"),
        "proyecto": col("proyecto relacionado", "proyecto"),
        "estado": col("estado", "status"),
        "fecha": col("fecha"),
        "prioridad": col("prioridad"),
        "comentarios": col("comentarios", "comentario"),
    }
    registros = []
    for fila in filas:
        if not fila or _vacio(fila[idx["actividad"]]):
            continue
        registros.append({k: (fila[i] if i is not None and i < len(fila) else None) for k, i in idx.items()})
    return nombre, semana, registros


def a_tarjeta(r: dict, semana_hoja: int, anio: int, proyectos: dict[str, int]) -> dict:
    estado_excel = str(r["estado"] or "").strip().lower()
    estado = ESTADOS_EXCEL.get(estado_excel, "backlog")
    responsable = None if _vacio(r["nombre"]) else str(r["nombre"]).strip()
    proyecto = None if _vacio(r["proyecto"]) else str(r["proyecto"]).strip()
    fecha = r["fecha"]
    if isinstance(fecha, datetime):
        fecha = fecha.date()
    elif isinstance(fecha, str):
        try:
            fecha = datetime.strptime(fecha.strip(), "%d/%m/%Y").date()
        except ValueError:
            fecha = None
    semana_origen = _semana(r["semana"]) or semana_hoja
    desde = datetime.combine(date.fromisocalendar(anio, semana_origen, 1), time(9), timezone.utc)
    t = {
        "titulo": str(r["actividad"]).strip(),
        "responsable": responsable,
        "proyecto": proyecto,
        "proyecto_id": proyectos.get(proyecto) if proyecto else None,
        "clase": "fecha_fija" if isinstance(fecha, date) else "estandar",
        "estado": estado,
        "fecha_compromiso": fecha.isoformat() if isinstance(fecha, date) else None,
        "comentario": None if _vacio(r["comentarios"]) else str(r["comentarios"]).strip(),
        "en_pausa": estado_excel == "en pausa",
        "semanas_en_excel": max(0, semana_hoja - semana_origen) + 1,
        "desde": desde,
    }
    # Lo pendiente que no cumple la Definición de Listo empieza en Backlog. Lo que
    # ya está en curso se respeta; sus huecos salen en el índice de higiene.
    if estado == "listo" and faltas_listo({**t, "proyecto_id": t["proyecto_id"] or proyecto}):
        t["estado"] = "backlog"
    return t


def diagnostico(nombre_hoja: str, registros: list[dict], tarjetas: list[dict]) -> str:
    total = len(registros)
    if total == 0:
        return "La hoja está vacía."
    pct = lambda n: f"{100 * n / total:.0f} %"  # noqa: E731
    estados = Counter(str(r["estado"] or "(vacío)").strip() for r in registros)
    sin_proyecto = sum(_vacio(r["proyecto"]) for r in registros)
    sin_fecha = sum(_vacio(r["fecha"]) for r in registros)
    sin_prioridad = sum(_vacio(r["prioridad"]) for r in registros)
    abiertas = [t for t in tarjetas if t["estado"] != "hecho"]
    viejas = [t for t in abiertas if t["semanas_en_excel"] >= 3]
    sin_empezar_viejas = [
        t for t, r in zip(tarjetas, registros)
        if str(r["estado"] or "").strip().lower() == "por iniciar" and t["semanas_en_excel"] >= 3
    ]
    compartidas = [t for t in tarjetas if t["responsable"] and re.search(r"[/,]| y ", t["responsable"])]
    carga = Counter(t["responsable"] or "(sin responsable)" for t in abiertas)

    lineas = [
        f"# Diagnóstico de {nombre_hoja}",
        "",
        f"- Actividades en la hoja: **{total}**",
        "- Por estado: " + ", ".join(f"{k} {v} ({pct(v)})" for k, v in estados.most_common()),
        f"- Sin proyecto: **{sin_proyecto} ({pct(sin_proyecto)})**",
        f"- Sin fecha de compromiso: **{sin_fecha} ({pct(sin_fecha)})**",
        f"- Sin prioridad: **{sin_prioridad} ({pct(sin_prioridad)})**",
        f"- Abiertas que arrastran 3 semanas o más: **{len(viejas)}**",
        f"- «Por iniciar» desde hace 3 semanas o más: **{len(sin_empezar_viejas)}** "
        "(candidatas a Backlog o descarte)",
        f"- Con responsable compartido: **{len(compartidas)}**",
        "",
        "## Carga abierta por persona (límite del marco: 3 en curso + 2 en validación)",
        "",
        *[f"- {p}: {n}{'  ⚠' if n > 5 else ''}" for p, n in carga.most_common()],
        "",
        "## Qué entraría en el tablero",
        "",
        *[f"- {k}: {v}" for k, v in Counter(t["estado"] for t in abiertas).most_common()],
        f"- Bloqueadas (venían «En pausa»): {sum(t['en_pausa'] for t in abiertas)}",
    ]
    return "\n".join(lineas)


def cargar(db: str, tarjetas: list[dict]) -> int:
    con = sqlite3.connect(db)
    con.executescript(ESQUEMA)
    ahora = datetime.now(timezone.utc).isoformat(timespec="seconds")
    proyectos = {}
    for nombre in sorted({t["proyecto"] for t in tarjetas if t["proyecto"]} | {"Interno"}):
        con.execute("INSERT OR IGNORE INTO flujo_proyectos (nombre) VALUES (?)", (nombre,))
        proyectos[nombre] = con.execute(
            "SELECT id FROM flujo_proyectos WHERE nombre=?", (nombre,)).fetchone()[0]
    n = 0
    for t in tarjetas:
        if t["estado"] == "hecho":
            continue
        desde = t["desde"].isoformat(timespec="seconds")
        cur = con.execute(
            "INSERT INTO flujo_tarjetas (titulo, responsable, proyecto_id, clase, estado, "
            "fecha_compromiso, bloqueada, bloqueo_motivo, bloqueo_depende_de, bloqueo_desde, "
            "creada_en, lista_en, iniciada_en, actualizada_en) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                t["titulo"], t["responsable"], proyectos.get(t["proyecto"]), t["clase"], t["estado"],
                t["fecha_compromiso"], int(t["en_pausa"]),
                (t["comentario"] or "Venía «En pausa» del Excel") if t["en_pausa"] else None,
                "Por definir" if t["en_pausa"] else None,
                desde if t["en_pausa"] else None,
                desde,
                desde if t["estado"] == "listo" else None,
                desde if t["estado"] == "en_curso" else None,
                ahora,
            ),
        )
        con.execute(
            "INSERT INTO flujo_eventos (tarjeta_id, momento, usuario, tipo, hacia, detalle) "
            "VALUES (?,?,?,?,?,?)",
            (cur.lastrowid, ahora, "importador", "importada", t["estado"], t["comentario"]),
        )
        n += 1
    con.commit()
    con.close()
    return n


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("excel")
    p.add_argument("--db", help="Base de datos de Flujo donde cargar las tarjetas abiertas")
    p.add_argument("--anio", type=int, default=date.today().year)
    p.add_argument("--salida", help="Guardar el diagnóstico en un fichero Markdown")
    args = p.parse_args(argv)

    nombre, semana, registros = leer_hoja_actual(args.excel)
    tarjetas = [a_tarjeta(r, semana, args.anio, {}) for r in registros]
    informe = diagnostico(nombre, registros, tarjetas)
    print(informe)
    if args.salida:
        with open(args.salida, "w", encoding="utf-8") as f:
            f.write(informe + "\n")
    if args.db:
        print(f"\nCargadas {cargar(args.db, tarjetas)} tarjetas abiertas en {args.db}.")


if __name__ == "__main__":
    main()

"""Prueba del importador con un Excel sintético (sin datos reales)."""

import sqlite3

from openpyxl import Workbook

from app import importar_seguimiento_cw as imp

CABECERA = ["Semana", "Nombre", "Área / Rol", "Actividad", "Proyecto Relacionado", "Estado",
            "% Avance", "Fecha", "Prioridad", "Comentarios", "Fecha Fin"]


def excel(tmp_path):
    wb = Workbook()
    wb.active.title = "CW39"
    wb["CW39"].append(CABECERA)
    hoja = wb.create_sheet("CW40")
    hoja.append(CABECERA)
    hoja.append(["CW 37", "Ana", "Ing.", "Elaborar protocolo", "Proyecto A", "Por Iniciar", "-", "-", "-", "-", None])
    hoja.append(["CW 40", "Luis", "Ing.", "Migrar servidor", "-", "En proceso", "40%", "-", "-", "-", None])
    hoja.append(["CW 39", "Ana", "Ing.", "Revisar pruebas", "Proyecto A", "En pausa", "-", "07/10/2026", "-", "Esperando a Comercial", None])
    hoja.append(["CW 40", "Eva / Luis", "Ing.", "Reparar pantallas", "Proyecto B", "Terminado", "100%", "-", "-", "-", None])
    oculta = wb.create_sheet("CW41")
    oculta.sheet_state = "hidden"
    ruta = tmp_path / "seguimiento.xlsx"
    wb.save(ruta)
    return str(ruta)


def test_diagnostico_y_carga(tmp_path, capsys):
    db = str(tmp_path / "flujo.db")
    imp.main([excel(tmp_path), "--db", db, "--anio", "2026"])
    salida = capsys.readouterr().out
    assert "Diagnóstico de CW40" in salida  # la hoja oculta CW41 se ignora
    assert "Sin proyecto: **1 (25 %)**" in salida
    assert "«Por iniciar» desde hace 3 semanas o más: **1**" in salida
    assert "Con responsable compartido: **1**" in salida

    con = sqlite3.connect(db)
    filas = con.execute("SELECT titulo, estado, bloqueada, fecha_compromiso FROM flujo_tarjetas ORDER BY id").fetchall()
    assert filas == [
        ("Elaborar protocolo", "listo", 0, None),
        ("Migrar servidor", "en_curso", 0, None),  # en curso se respeta aunque falte proyecto
        ("Revisar pruebas", "en_curso", 1, "2026-10-07"),
    ]
